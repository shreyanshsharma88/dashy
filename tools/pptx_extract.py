#!/usr/bin/env python3
"""dashy pptx extractor: slide text + speaker notes + embedded images.
Usage: pptx_extract.py <file.pptx> <out-prefix>
Writes <out-prefix>.txt and <out-prefix>-img<N>.<ext>, prints JSON summary.
Requires: python-pptx, pillow (dashy bootstraps ~/.dashy-tools venv).
"""
import sys, os, json, hashlib

def shape_texts(shape, out):
    try:
        if shape.shape_type == 6:  # group shape: recurse
            for s in shape.shapes:
                shape_texts(s, out)
            return
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs).strip()
                if t:
                    out.append(t)
        if shape.has_table:
            for row in shape.table.rows:
                line = " | ".join(c.text.strip() for c in row.cells).strip(" |")
                if line:
                    out.append(line)
    except Exception:
        pass

def main():
    from pptx import Presentation
    notes_only = "--notes-only" in sys.argv
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    src, pre = args[0], (args[1] if len(args) > 1 else None)
    prs = Presentation(src)
    seen = set()
    nimg = 0
    chunks = []
    notes_all = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            shape_texts(shape, texts)
            try:
                if not notes_only and shape.shape_type == 13 and shape.image.blob:  # picture
                    h = hashlib.md5(shape.image.blob).hexdigest()
                    if h not in seen:
                        seen.add(h)
                        ext = shape.image.ext or "png"
                        nimg += 1
                        with open("%s-img%d.%s" % (pre, nimg, ext), "wb") as f:
                            f.write(shape.image.blob)
            except Exception:
                pass
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.placeholders[1].text.strip()
        except Exception:
            pass
        chunks.append("--- SLIDE %d ---\n%s%s" % (i, "\n".join(texts), ("\n[Notes] " + notes) if notes else ""))
        notes_all.append(notes)
    if notes_only:
        print("\n\n".join("SLIDE %d notes: %s" % (i, n) for i, n in enumerate(notes_all, 1) if n) or "(no speaker notes)")
        return
    with open(pre + ".txt", "w", encoding="utf-8") as f:
        f.write("\n\n".join(chunks) + "\n")
    print(json.dumps({"slides": len(prs.slides), "images": nimg,
                      "lines": sum(c.count("\n") for c in chunks) + len(chunks)}))

main()
