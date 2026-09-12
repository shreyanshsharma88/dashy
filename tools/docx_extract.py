#!/usr/bin/env python3
"""dashy docx extractor: paragraphs + tables to stdout (or file).
Usage: docx_extract.py <file.docx> [out.txt]
"""
import sys

def main():
    from docx import Document
    src = sys.argv[1]
    doc = Document(src)
    parts = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for tb in doc.tables:
        for row in tb.rows:
            line = " | ".join(c.text.strip() for c in row.cells).strip(" |")
            if line:
                parts.append(line)
    text = "\n".join(parts) + "\n"
    if len(sys.argv) > 2:
        with open(sys.argv[2], "w", encoding="utf-8") as f:
            f.write(text)
    else:
        sys.stdout.write(text)

main()
